"""文本抽取与分块工具。

支持格式：PDF、TXT/Markdown、Word（.docx）。抽取后统一 ``normalize_text``，
再按字符窗口 ``chunk_text`` 切块，优先在段落/句号等边界处断开，兼顾中英文。

常量 ``SUPPORTED_EXTENSIONS`` 与上传校验、前端 ``accept`` 需保持一致。
"""
from __future__ import annotations

import re
from pathlib import Path
from typing import List

from docx import Document as DocxDocument
from pypdf import PdfReader


SUPPORTED_EXTENSIONS = {".pdf", ".txt", ".md", ".markdown", ".docx"}

# Word 的 Content-Type 较长，用子串匹配更稳妥
DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


def extract_text(file_path: str, content_type: str | None = None) -> str:
    """根据扩展名或 MIME 从磁盘文件抽取纯文本。

    Args:
        file_path: 磁盘绝对路径。
        content_type: 浏览器/客户端上报的 MIME，辅助判断 PDF/Word。

    Returns:
        合并后的纯文本字符串（可能为空）。

    Raises:
        ValueError: 不支持的扩展名。
    """
    p = Path(file_path)
    ext = p.suffix.lower()
    ct = (content_type or "").lower()

    if ext == ".pdf" or "pdf" in ct:
        return _extract_pdf(p)
    if ext == ".docx" or "wordprocessingml" in ct or ct == DOCX_MIME:
        return _extract_docx(p)
    if ext in {".txt", ".md", ".markdown"} or "text" in ct:
        return p.read_text(encoding="utf-8", errors="ignore")
    raise ValueError(f"Unsupported file extension: {ext}")


def _extract_pdf(p: Path) -> str:
    """逐页抽取 PDF 文本，单页失败则跳过该页。"""
    parts: List[str] = []
    with p.open("rb") as f:
        reader = PdfReader(f)
        for page in reader.pages:
            try:
                parts.append(page.extract_text() or "")
            except Exception:
                continue
    return "\n".join(parts)


def _extract_docx(p: Path) -> str:
    """抽取 Word 段落与表格单元格文本；表格行内用 `` | `` 连接。"""
    doc = DocxDocument(str(p))
    parts: List[str] = []
    for para in doc.paragraphs:
        text = (para.text or "").strip()
        if text:
            parts.append(text)
    for table in doc.tables:
        for row in table.rows:
            cells = [
                (cell.text or "").strip()
                for cell in row.cells
                if (cell.text or "").strip()
            ]
            if cells:
                parts.append(" | ".join(cells))
    return "\n".join(parts)


_WS_RE = re.compile(r"[ \t]+")
_MULTI_NEWLINE_RE = re.compile(r"\n{3,}")


def normalize_text(text: str) -> str:
    """统一换行、压缩连续空白与过多空行，便于后续切块。"""
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = _WS_RE.sub(" ", text)
    text = _MULTI_NEWLINE_RE.sub("\n\n", text)
    return text.strip()


def chunk_text(text: str, chunk_size: int = 800, overlap: int = 120) -> List[str]:
    """按字符长度滑动窗口切块，窗口尾部优先在标点/换行处断开。

    Args:
        text: 已规范化全文。
        chunk_size: 单块最大字符数。
        overlap: 相邻块重叠字符数（不超过 chunk_size 的一半）。

    Returns:
        非空文本块列表，顺序与原文一致。

    Raises:
        ValueError: ``chunk_size`` 非正。
    """
    if chunk_size <= 0:
        raise ValueError("chunk_size must be positive")
    overlap = max(0, min(overlap, chunk_size // 2))

    text = normalize_text(text)
    if not text:
        return []

    if len(text) <= chunk_size:
        return [text]

    chunks: List[str] = []
    start = 0
    n = len(text)
    while start < n:
        end = min(start + chunk_size, n)
        # 在窗口后 30% 区域内从后向前找第一个分隔符，减少硬切句中
        if end < n:
            window_start = start + int(chunk_size * 0.7)
            slice_ = text[window_start:end]
            for sep in ["\n\n", "\n", "。", ". ", "！", "？", "!", "?", "; ", "；"]:
                idx = slice_.rfind(sep)
                if idx != -1:
                    end = window_start + idx + len(sep)
                    break
        chunk = text[start:end].strip()
        if chunk:
            chunks.append(chunk)
        if end >= n:
            break
        start = max(0, end - overlap)
    return chunks
